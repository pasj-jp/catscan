import io
import re
from typing import Any, Dict, List, Optional, Tuple, Union
from docx import Document

from ..base import BasePaperChecker
from ..jacow.score import score
from ..jacow.margins import check_sections
from ..jacow.page import check_tracking_on, check_title_case
from ..jacow.styles import get_style_details
from ..engine import ElementContext, evaluate_section
from ..rules import load_rules
from ..rules.schema import WordRules



def safe_style_details(p) -> Dict[str, Any]:
    try:
        st = get_style_details(p)
        st.pop("p", None)
        st.pop("url", None)
        return {
            k: v for k, v in st.items()
            if isinstance(v, (str, int, float, bool)) or v is None
        }
    except Exception:
        return {}


def _get_pasj_rules() -> WordRules:
    return load_rules("pasj_word", WordRules)


# Backward-compatible accessors for code/tests expecting dictionary access
PASJ_WORD_RULES = {
    k: {"rules": v.rules, "extra_rules": v.extra_rules}
    for k, v in _get_pasj_rules().sections.items()
}
PASJ_REQUIRED_STYLES = _get_pasj_rules().required_styles



def _ends_with_period_or_note(text: str) -> bool:
    """
    Check if a figure caption ends with a period '.', or a period followed by a parenthetical
    credit/note such as '. (Publishing approved by JACoW)'.
    """
    t = text.strip()
    if t.endswith("."):
        return True
    if re.search(r'\.\s*\([^)]+\)\s*$', t):
        return True
    return False


class PasjWordChecker(BasePaperChecker):
    """
    PASJ Japanese Word document validator.
    Inspects Microsoft Word (.docx) documents for Japanese Accelerator Society (PASJ)
    2026 guidelines:
    - Bilingual titles (和文タイトル 14pt ゴシック / 英文タイトル 12pt Times 太字・大文字)
    - Bilingual authors (和文・英文著者 10pt、半角カンマ+スペース区切り)
    - English abstract (見出し 12pt 斜体 / 本文 10pt 英語のみ)
    - Japanese punctuation rules (「、」「。」 vs 「，」「．」)
    - English figure & table captions (in both body paragraphs and layout tables)
    - Headings, A4 margins, 2-column layout, PASJ styles, and references
    """

    def validate(
        self,
        file_source: Union[str, bytes, io.BytesIO],
        filename: str = "document.docx",
        conference_id: Optional[Union[str, int]] = None,
    ) -> Dict[str, Any]:
        if isinstance(file_source, bytes):
            file_obj = io.BytesIO(file_source)
        elif isinstance(file_source, str):
            with open(file_source, "rb") as f:
                file_obj = io.BytesIO(f.read())
        elif isinstance(file_source, io.BytesIO):
            file_obj = file_source
        else:
            return {"error": "Invalid file source provided."}

        try:
            doc = Document(docx=file_obj)
        except Exception as e:
            return {"error": f"Failed to open Word document: {str(e)}"}

        try:
            if check_tracking_on(doc):
                return {
                    "error": "Cannot process file: Track Changes (revision tracking) is turned on.\n"
                             "Please accept/reject all changes and turn tracking off before submitting."
                }
        except Exception:
            pass

        summary = self._analyze_document(doc)
        scores = score(summary)

        meta = {}
        metadata = doc.core_properties
        if metadata:
            meta = {
                "author": metadata.author,
                "revision": metadata.revision,
                "version": metadata.version,
                "language": metadata.language,
            }
            if metadata.created:
                meta["created"] = metadata.created.strftime("%d/%m/%Y %H:%M")
            if metadata.modified:
                meta["modified"] = metadata.modified.strftime("%d/%m/%Y %H:%M")

        return {
            "scores": scores,
            "summary": summary,
            "metadata": meta,
            "filename": filename,
            "detected_template": "pasj",
            "template_name": "PASJ 和文原稿",
        }

    def _analyze_document(self, doc: Document) -> Dict[str, Any]:
        jtitle_p = None
        etitle_p = None
        jauthor_p = []
        eauthor_p = []
        abstract_heading_p = None
        abstract_p = []
        body_p = []
        heading_p = []
        ref_heading_p = None
        ref_p = []
        fig_cap_p = []
        tbl_cap_p = []

        seen_abstract = False
        seen_refs = False
        seen_caption_elements = set()

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else ""

            if "参考文献見出し" in style_name or text == "参考文献" or text.lower() == "references":
                seen_refs = True
                ref_heading_p = p
                continue

            if seen_refs:
                ref_p.append(p)
                continue

            if "Abstract" in style_name or text.lower() == "abstract":
                seen_abstract = True
                abstract_heading_p = p
                continue

            if not seen_abstract:
                if "和文タイトル" in style_name or "011" in style_name:
                    if jtitle_p is None:
                        jtitle_p = p
                        continue
                if "英文タイトル" in style_name or "012" in style_name:
                    if etitle_p is None:
                        etitle_p = p
                        continue
                if "著者" in style_name or "020" in style_name or "author" in style_name.lower():
                    if re.search(r'[\u3040-\u30ff\u4e00-\u9faf]', text):
                        jauthor_p.append(p)
                    else:
                        eauthor_p.append(p)
                    continue
                if jtitle_p is None and re.search(r'[\u3040-\u30ff\u4e00-\u9faf]', text):
                    jtitle_p = p
                elif etitle_p is None and jtitle_p is not None and not re.search(r'[\u3040-\u30ff\u4e00-\u9faf]', text) and not ("著者" in style_name or "author" in style_name.lower()):
                    etitle_p = p
                elif re.search(r'[\u3040-\u30ff\u4e00-\u9faf]', text):
                    jauthor_p.append(p)
                else:
                    eauthor_p.append(p)
            else:
                if len(abstract_p) == 0 and not ("見出し" in style_name or "heading" in style_name.lower()):
                    abstract_p.append(p)
                    continue

                if "謝辞見出し" in style_name or "063" in style_name or text == "謝辞":
                    continue
                elif "章見出し" in style_name or "節見出し" in style_name or "項見出し" in style_name or re.match(r'^\d+(?:\.\d+)*\.?\s', text):
                    if not ("箇条書き" in style_name or "07" in style_name):
                        heading_p.append(p)
                    else:
                        body_p.append(p)
                elif ("図タイトル" in style_name or re.match(r'^(?:Figure|Fig\.)[\s\xa0]*\d+\s*[:.]', text, re.IGNORECASE)) and not re.match(r'^\([a-z0-9]+\)', text, re.IGNORECASE):
                    fig_cap_p.append(p)
                    seen_caption_elements.add(id(p._p))
                elif "表タイトル" in style_name or re.match(r'^Table[\s\xa0]*\d+\s*[:.]', text, re.IGNORECASE):
                    tbl_cap_p.append(p)
                    seen_caption_elements.add(id(p._p))
                else:
                    body_p.append(p)

        # Also inspect tables for Figure captions and Table titles (standard in PASJ Word template)
        for t in doc.tables:
            seen_cells = set()
            for r in t.rows:
                for c in r.cells:
                    if id(c._tc) in seen_cells:
                        continue
                    seen_cells.add(id(c._tc))
                    for p in c.paragraphs:
                        if id(p._p) in seen_caption_elements:
                            continue
                        text = p.text.strip()
                        if not text:
                            continue
                        style_name = p.style.name if p.style else ""
                        # Ignore subfigure labels like "(a) PNG image"
                        if re.match(r'^\([a-z0-9]+\)', text, re.IGNORECASE):
                            continue
                        if "図タイトル" in style_name or re.match(r'^(?:Figure|Fig\.)[\s\xa0]*\d+\s*[:.]', text, re.IGNORECASE):
                            fig_cap_p.append(p)
                            seen_caption_elements.add(id(p._p))
                        elif "表タイトル" in style_name or re.match(r'^Table[\s\xa0]*\d+\s*[:.]', text, re.IGNORECASE):
                            tbl_cap_p.append(p)
                            seen_caption_elements.add(id(p._p))

        summary = {}

        # 1. Title summary (evaluated via declarative rules engine)
        rules = load_rules("pasj_word", WordRules)
        elements_title = {
            "jtitle": ElementContext(
                text=jtitle_p.text,
                style_name=jtitle_p.style.name if jtitle_p.style else "",
                style_details=safe_style_details(jtitle_p),
                paragraph=jtitle_p,
            ) if jtitle_p else None,
            "etitle": ElementContext(
                text=etitle_p.text,
                style_name=etitle_p.style.name if etitle_p.style else "",
                style_details=safe_style_details(etitle_p),
                paragraph=etitle_p,
            ) if etitle_p else None,
        }
        summary["Title"] = evaluate_section(rules.sections["Title"], elements_title)

        # 2. Authors summary (evaluated via declarative rules engine)
        elements_authors = {
            "jauthor": [
                ElementContext(
                    text=p.text,
                    style_name=p.style.name if p.style else "",
                    style_details=safe_style_details(p),
                    paragraph=p,
                )
                for p in jauthor_p
            ],
            "eauthor": [
                ElementContext(
                    text=p.text,
                    style_name=p.style.name if p.style else "",
                    style_details=safe_style_details(p),
                    paragraph=p,
                )
                for p in eauthor_p
            ],
        }
        summary["Authors"] = evaluate_section(rules.sections["Authors"], elements_authors)


        # 3. Abstract summary
        elements_abstract = {
            "heading": ElementContext(
                text=abstract_heading_p.text,
                style_name=abstract_heading_p.style.name if abstract_heading_p.style else "",
                style_details=safe_style_details(abstract_heading_p),
                paragraph=abstract_heading_p,
            ) if abstract_heading_p else None,
            "body": [
                ElementContext(
                    text=p.text,
                    style_name=p.style.name if p.style else "",
                    style_details=safe_style_details(p),
                    paragraph=p,
                )
                for p in abstract_p
            ],
        }
        summary["Abstract"] = evaluate_section(rules.sections["Abstract"], elements_abstract)

        # 4. Headings summary
        elements_headings = {
            "heading": [
                ElementContext(
                    text=p.text,
                    style_name=p.style.name if p.style else "",
                    style_details=safe_style_details(p),
                    paragraph=p,
                )
                for p in heading_p
            ]
        }
        summary["Headings"] = evaluate_section(rules.sections["Headings"], elements_headings)

        # 5. Paragraphs / Punctuation summary
        fullwidth_commas = 0
        fullwidth_periods = 0
        correct_punct = 0
        comma_examples: List[str] = []
        period_examples: List[str] = []

        for p in body_p:
            t = p.text.strip()
            c_cnt = t.count("，")
            p_cnt = t.count("．")
            fullwidth_commas += c_cnt
            fullwidth_periods += p_cnt
            correct_punct += t.count("、") + t.count("。")
            if c_cnt > 0 and len(comma_examples) < 2:
                comma_examples.append(t[:45] + ("..." if len(t) > 45 else ""))
            if p_cnt > 0 and len(period_examples) < 2:
                period_examples.append(t[:45] + ("..." if len(t) > 45 else ""))

        punct_ok = (fullwidth_commas == 0 and fullwidth_periods == 0)
        comma_status = "問題なし (0 箇所)"
        if fullwidth_commas > 0:
            ex_str = f" — 例: 「{comma_examples[0]}」" if comma_examples else ""
            comma_status = f"和文原稿では使用禁止 ({fullwidth_commas} 箇所検出{ex_str})。「、」に置換してください"

        period_status = "問題なし (0 箇所)"
        if fullwidth_periods > 0:
            ex_str = f" — 例: 「{period_examples[0]}」" if period_examples else ""
            period_status = f"和文原稿では使用禁止 ({fullwidth_periods} 箇所検出{ex_str})。「。」に置換してください"

        para_details = [
            {
                "item": "全角句読点（「、」「。」）",
                "count": f"{correct_punct} 箇所",
                "status": "推奨書式（公用文・PASJ規定準拠）",
                "punct_ok": True,
            },
            {
                "item": "全角カンマ（「，」）",
                "count": f"{fullwidth_commas} 箇所",
                "status": comma_status,
                "punct_ok": fullwidth_commas == 0,
            },
            {
                "item": "全角ピリオド（「．」）",
                "count": f"{fullwidth_periods} 箇所",
                "status": period_status,
                "punct_ok": fullwidth_periods == 0,
            },
        ]

        summary["Paragraphs"] = {
            "title": rules.sections["Paragraphs"].title,
            "ok": punct_ok,
            "details": para_details,
            "extra_info": rules.sections["Paragraphs"].extra_info.to_output(),
            "rules": dict(rules.sections["Paragraphs"].rules),
            "extra_rules": list(rules.sections["Paragraphs"].extra_rules),
            "anchor": rules.sections["Paragraphs"].anchor,
        }

        # 6. Figures summary
        elements_figures = {
            "figure": [
                ElementContext(
                    text=p.text,
                    style_name=p.style.name if p.style else "",
                    style_details=safe_style_details(p),
                    paragraph=p,
                )
                for p in fig_cap_p
            ]
        }
        summary["Figures"] = evaluate_section(rules.sections["Figures"], elements_figures)

        # 7. Tables summary
        elements_tables = {
            "table": [
                ElementContext(
                    text=p.text,
                    style_name=p.style.name if p.style else "",
                    style_details=safe_style_details(p),
                    paragraph=p,
                )
                for p in tbl_cap_p
            ]
        }
        summary["Tables"] = evaluate_section(rules.sections["Tables"], elements_tables)

        # 8. References summary
        elements_refs = {
            "reference": [
                ElementContext(
                    text=p.text,
                    style_name=p.style.name if p.style else "",
                    style_details=safe_style_details(p),
                    paragraph=p,
                )
                for p in ref_p
            ]
        }
        summary["References"] = evaluate_section(rules.sections["References"], elements_refs)

        # 9. Margins summary (A4 layout check: Top 37mm, Bottom 19mm, Left 20mm, Right 20mm, 2 cols)
        sections = check_sections(doc)
        margins_ok = all(s["margins_ok"] and s["col_ok"] for s in sections)
        summary["Margins"] = {
            "title": rules.sections["Margins"].title,
            "extra_rules": list(rules.sections["Margins"].extra_rules),
            "extra_info": rules.sections["Margins"].extra_info.to_output(),
            "ok": margins_ok,
            "message": rules.sections["Margins"].message or "Margins",
            "details": sections,
            "anchor": rules.sections["Margins"].anchor,
        }

        # 10. Styles summary
        doc_styles = [s.name for s in doc.styles if s.name]
        style_details = [
            {"style": s, "style_ok": s in doc_styles}
            for s in rules.required_styles
        ]
        summary["Styles"] = {
            "title": rules.sections["Styles"].title,
            "ok": all(d.get("style_ok") for d in style_details),
            "details": style_details,
            "extra_info": rules.sections["Styles"].extra_info.to_output(),
            "rules": dict(rules.sections["Styles"].rules),
            "extra_rules": list(rules.sections["Styles"].extra_rules),
            "anchor": rules.sections["Styles"].anchor,
        }


        return summary
